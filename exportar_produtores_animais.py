# -*- coding: utf-8 -*-
"""
Gera a planilha "produtores_animais.xlsx" com todos os produtores e os
animais (receptoras) vinculados a cada um, lendo direto do Supabase.

COMO USAR
---------
1. Coloque este arquivo na pasta do projeto (do lado do database.py e do .env).
2. Rode:  python exportar_produtores_animais.py
3. A planilha sai na mesma pasta: produtores_animais.xlsx

Abas geradas:
  - "Animais por Produtor": uma linha por animal vinculado (produtor sem
    animal aparece com os campos do animal em branco).
  - "Resumo": uma linha por produtor com a quantidade de animais.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SAIDA = Path(__file__).parent / "produtores_animais.xlsx"

SQL = """
SELECT p.id AS produtor_id,
       p.nome_produtor,
       p.cpf,
       p.municipio,
       p.nome_propriedade,
       p.tecnico_responsavel,
       a.brinco_faec,
       a.brinco_fazenda,
       a.peso
FROM produtores p
LEFT JOIN animais a ON a.produtor_id = p.id
ORDER BY p.nome_produtor, a.brinco_faec
"""

FONTE = "Arial"
COR_CABECALHO = "1F4E78"
BORDA = Side(style="thin", color="BFBFBF")


def buscar_linhas():
    from database import get_connection

    conn = get_connection()
    try:
        return [dict(r) for r in conn.execute(SQL).fetchall()]
    finally:
        conn.close()


def _peso(valor):
    """O peso vem como texto (ex: '320,5'). Vira número quando der."""
    if valor is None or str(valor).strip() == "":
        return None
    try:
        return float(str(valor).strip().replace(",", "."))
    except ValueError:
        return str(valor).strip()


def _cabecalho(ws, titulos):
    for col, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=1, column=col, value=titulo)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def _largura(ws, larguras):
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura


def gerar_planilha(linhas, caminho=SAIDA):
    wb = Workbook()

    # ---------------- Aba 1: animais por produtor ----------------
    ws = wb.active
    ws.title = "Animais por Produtor"
    _cabecalho(ws, ["ID", "Produtor", "CPF", "Município", "Propriedade",
                    "Técnico", "Brinco FAEC", "Brinco Fazenda", "Peso (kg)"])

    for i, r in enumerate(linhas, start=2):
        valores = [
            r["produtor_id"], r["nome_produtor"], r["cpf"], r["municipio"],
            r["nome_propriedade"], r["tecnico_responsavel"],
            r["brinco_faec"], r["brinco_fazenda"], _peso(r["peso"]),
        ]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.font = Font(name=FONTE, size=10)
            c.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)
        # CPF e brincos ficam como texto (preserva zeros à esquerda)
        for col in (3, 7, 8):
            ws.cell(row=i, column=col).number_format = "@"
        ws.cell(row=i, column=9).number_format = "#,##0.0"

    ultima = len(linhas) + 1
    _largura(ws, [7, 38, 16, 20, 28, 18, 14, 16, 11])
    if linhas:
        ws.auto_filter.ref = f"A1:I{ultima}"

    # ---------------- Aba 2: resumo por produtor ----------------
    rs = wb.create_sheet("Resumo")
    _cabecalho(rs, ["ID", "Produtor", "Município", "Propriedade", "Qtd. de animais"])

    vistos, produtores = set(), []
    for r in linhas:
        if r["produtor_id"] not in vistos:
            vistos.add(r["produtor_id"])
            produtores.append(r)

    for i, r in enumerate(produtores, start=2):
        for col, v in enumerate(
            [r["produtor_id"], r["nome_produtor"], r["municipio"], r["nome_propriedade"]],
            start=1,
        ):
            c = rs.cell(row=i, column=col, value=v)
            c.font = Font(name=FONTE, size=10)
            c.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)
        c = rs.cell(
            row=i, column=5,
            value=(f"=COUNTIFS('Animais por Produtor'!$A$2:$A${ultima},A{i},"
                   f"'Animais por Produtor'!$G$2:$G${ultima},\"<>\")"),
        )
        c.font = Font(name=FONTE, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)

    fim = len(produtores) + 1
    total = fim + 1
    rs.cell(row=total, column=4, value="Total").font = Font(name=FONTE, bold=True)
    c = rs.cell(row=total, column=5, value=f"=SUM(E2:E{fim})" if produtores else 0)
    c.font = Font(name=FONTE, bold=True)
    c.alignment = Alignment(horizontal="center")

    rs.cell(row=total + 2, column=2,
            value=f"Gerado em {datetime.now():%d/%m/%Y %H:%M} a partir do Supabase "
                  "(tabelas produtores e animais).").font = Font(
        name=FONTE, italic=True, size=9, color="7F7F7F")

    _largura(rs, [7, 38, 20, 28, 16])
    if produtores:
        rs.auto_filter.ref = f"A1:E{fim}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(caminho)
    return len(produtores), sum(1 for r in linhas if r["brinco_faec"])


if __name__ == "__main__":
    dados = buscar_linhas()
    n_prod, n_animais = gerar_planilha(dados)
    print(f"Pronto: {n_prod} produtores, {n_animais} animais vinculados.")
    print(f"Arquivo: {SAIDA}")
