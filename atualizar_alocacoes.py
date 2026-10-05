# -*- coding: utf-8 -*-
"""
Atualiza QUEM e o dono de cada animal (brinco FAEC -> produtor) a partir de
uma planilha, SEM MEXER NAS FOTOS que ja estao cadastradas.

POR QUE AS FOTOS FICAM SEGURAS
------------------------------
As fotos (foto_1, foto_2, foto_3) moram na MESMA linha do animal, na tabela
"FIV".animais. Trocar o dono do animal e so mudar duas colunas dessa linha
(produtor_id e status). Este script:

  * faz UPDATE somente de produtor_id, status e atualizado_em;
  * NUNCA apaga e NUNCA recria um animal (apagar+recriar e o que faria a foto
    se perder);
  * NUNCA mexe nas colunas foto_1/foto_2/foto_3 nem no Storage;
  * antes de gravar, guarda um backup completo da tabela em CSV;
  * confere, dentro da mesma transacao, que as fotos de TODOS os animais
    continuam identicas. Se qualquer foto mudar, desfaz tudo (ROLLBACK).

COMO USAR (na pasta do projeto, com o .env configurado)
--------------------------------------------------------
1) SIMULACAO (padrao - nao grava nada, so mostra o que mudaria):

       python atualizar_alocacoes.py animais_alocados_fats.xlsx

2) Se o relatorio estiver certo, grave de verdade:

       python atualizar_alocacoes.py animais_alocados_fats.xlsx --aplicar

OPCOES
------
  --aplicar            grava no banco (sem isso, e so simulacao)
  --liberar-ausentes   animais que estao alocados no banco mas NAO aparecem
                       na planilha voltam a "disponivel" (sem produtor).
                       Sem esta opcao eles nao sao tocados, so listados.
  --ignorar-limite     permite passar de 5 animais por produtor.

FORMATO DA PLANILHA
-------------------
Primeira linha com os titulos. Colunas usadas (as demais sao ignoradas):
    ID Produtor | Produtor | CPF | Brinco FAEC | Brinco Fazenda (recomendada)
- "ID Produtor" e o id do produtor no sistema (conferido pelo CPF).
- O brinco FAEC e o brinco da fazenda sao ATRELADOS: sao o mesmo animal (e as
  fotos dele). Quando a planilha traz o "Brinco Fazenda" (o export do sistema
  traz), a rotina confere que o par bate com o do banco; se nao bater, NAO
  grava, para o animal nao ir para o lugar errado. E essa coluna tambem
  resolve brinco FAEC repetido no banco (a fazenda diz qual e qual).
- Uma linha por animal. Linha com produtor mas SEM brinco = vaga em aberto
  (nao muda nada, so e avisada). Linhas totalmente vazias sao ignoradas.

O QUE IMPEDE A GRAVACAO (e precisa ser corrigido na planilha)
--------------------------------------------------------------
- o mesmo animal para dois produtores diferentes;
- brinco sem produtor na linha;
- produtor que nao existe no sistema, ou cujo CPF nao confere com o ID;
- brinco FAEC da planilha atrelado a OUTRO brinco da fazenda no banco;
- brinco repetido no banco sem a planilha dizer qual e (falta o Brinco Fazenda);
- produtor ficando com mais de 5 animais (a menos que use --ignorar-limite).
"""

import argparse
import csv
import hashlib
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from database import get_connection

LIMITE_POR_PRODUTOR = 5  # mesmo limite do sistema (app.py)
PASTA_SAIDA = Path(__file__).parent / "backups_alocacoes"


# ------------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------------
def so_digitos(valor) -> str:
    return re.sub(r"\D", "", str(valor or ""))


def sem_acento_maiusculo(texto) -> str:
    t = unicodedata.normalize("NFD", str(texto or ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", t).strip().upper()


def normalizar_brinco(valor) -> str:
    """'18', 18, 18.0 e '0018' viram todos '0018'. Brincos com letras so
    ganham espacos/maiusculas padronizados."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = re.sub(r"\s+", " ", str(valor).strip().upper())
    if texto.endswith(".0") and texto[:-2].isdigit():
        texto = texto[:-2]
    return texto.zfill(4) if texto.isdigit() else texto


def chave_fazenda(valor) -> str:
    """Brinco da fazenda para COMPARAR: so letras e numeros, maiusculo
    ('V 1821', 'v-1821' e 'V1821' sao o mesmo)."""
    return re.sub(r"[^A-Z0-9]", "", sem_acento_maiusculo(valor))


def _chave_coluna(texto) -> str:
    return re.sub(r"[^a-z0-9]", "", sem_acento_maiusculo(texto).lower())


def _texto_celula(valor):
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


# ------------------------------------------------------------------
# Leitura da planilha
# ------------------------------------------------------------------
def ler_planilha(caminho: Path):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = next(linhas, None)
    if not cabecalho:
        raise SystemExit("A planilha esta vazia.")

    destino = {"idprodutor": "id", "produtor": "nome", "cpf": "cpf", "brincofaec": "brinco",
               "brincofazenda": "fazenda"}
    colunas = {}
    for posicao, titulo in enumerate(cabecalho):
        campo = destino.get(_chave_coluna(titulo))
        if campo and campo not in colunas:
            colunas[campo] = posicao
    faltando = [t for t, c in (("ID Produtor", "id"), ("Brinco FAEC", "brinco")) if c not in colunas]
    if faltando:
        raise SystemExit(
            "Nao encontrei na primeira linha da planilha a(s) coluna(s): "
            + ", ".join(faltando) + "."
        )

    registros = []
    for n, linha in enumerate(linhas, start=2):
        def pega(campo):
            pos = colunas.get(campo)
            return _texto_celula(linha[pos]) if pos is not None and pos < len(linha) else None

        reg = {
            "linha": n,
            "id": pega("id"),
            "nome": pega("nome"),
            "cpf": pega("cpf"),
            "brinco": normalizar_brinco(pega("brinco")),
            "fazenda": pega("fazenda"),
        }
        if not any([reg["id"], reg["nome"], reg["cpf"], reg["brinco"]]):
            continue  # linha totalmente vazia
        if reg["id"]:
            try:
                reg["id"] = str(int(float(reg["id"])))
            except ValueError:
                pass
        registros.append(reg)
    wb.close()
    return registros, ("fazenda" in colunas)


# ------------------------------------------------------------------
# Montagem do plano (so leitura)
# ------------------------------------------------------------------
class Plano:
    def __init__(self):
        self.bloqueios = []
        self.avisos = []
        self.movimentos = []      # dicts: acao, animal_id, brinco, atual_id, novo_id, novo_status
        self.iguais = 0
        self.sem_cadastro = []    # brincos da planilha que nao existem no banco
        self.vagas = []           # produtores com linha sem brinco
        self.ausentes = []        # alocados no banco e fora da planilha
        self.nomes = {}           # produtor_id -> nome (para o relatorio)


def _pares_oficiais():
    """Lista oficial (brinco FAEC -> brinco da fazenda) usada na importacao inicial dos animais
    (importar_animais.py). Serve so para AVISAR se o par de um animal mexido foi alterado depois."""
    try:
        from importar_animais import ANIMAIS
    except Exception:  # noqa: BLE001 - se nao der para importar, apenas nao confere
        return {}
    return {normalizar_brinco(f): fz for f, fz in ANIMAIS}


def _qtd_fotos(animal):
    return sum(1 for c in ("foto_1", "foto_2", "foto_3") if animal.get(c))


def montar_plano(conn, registros, liberar_ausentes=False, ignorar_limite=False, tem_fazenda=False) -> Plano:
    plano = Plano()

    # 1) Problemas da propria planilha ---------------------------------
    # Cada animal = brinco FAEC + brinco da fazenda (atrelados). Sem a coluna da fazenda,
    # o brinco FAEC sozinho identifica o animal.
    def chave_animal(r):
        return (r["brinco"], chave_fazenda(r["fazenda"]) if tem_fazenda else "")

    por_animal = defaultdict(list)      # (faec, fazenda) -> [(id_planilha, linha, fazenda_original)]
    for r in registros:
        if r["brinco"] and not r["id"]:
            plano.bloqueios.append(f"Linha {r['linha']}: o brinco {r['brinco']} esta sem produtor.")
        elif r["id"] and not r["brinco"]:
            plano.vagas.append(r)
        elif r["id"] and r["brinco"]:
            por_animal[chave_animal(r)].append((r["id"], r["linha"], r["fazenda"]))

    for (brinco, fz), usos in sorted(por_animal.items()):
        ids = sorted({i for i, _, _ in usos})
        if len(ids) > 1:
            onde = "; ".join(f"produtor {i} (linha {l})" for i, l, _ in usos)
            extra = f" ({usos[0][2]})" if fz else ""
            plano.bloqueios.append(f"O animal de brinco {brinco}{extra} aparece para mais de um produtor: {onde}.")

    # o mesmo brinco FAEC em produtores diferentes so e aceito quando a planilha traz a fazenda
    # de cada um e elas sao diferentes (sao animais diferentes com o mesmo numero de brinco FAEC)
    faec_para_ids = defaultdict(set)
    faec_sem_fazenda = set()
    for (brinco, fz), usos in por_animal.items():
        for i, l, orig in usos:
            faec_para_ids[brinco].add(i)
            if tem_fazenda and not chave_fazenda(orig):
                faec_sem_fazenda.add(brinco)
    for brinco, ids in sorted(faec_para_ids.items()):
        if len(ids) > 1 and (not tem_fazenda or brinco in faec_sem_fazenda):
            onde = "; ".join(f"produtor {i}" for i in sorted(ids))
            plano.bloqueios.append(
                f"O brinco FAEC {brinco} aparece em mais de um produtor ({onde}) e a planilha nao diz o brinco da "
                f"fazenda de cada um. Inclua a coluna 'Brinco Fazenda' (o export do sistema traz) para o sistema saber qual e qual."
            )

    # 2) Produtores: ID da planilha -> produtor do banco ----------------
    produtores = conn.execute("SELECT id, nome_produtor, cpf FROM produtores").fetchall()
    por_id = {str(p["id"]): p for p in produtores}
    por_cpf = defaultdict(list)
    for p in produtores:
        if so_digitos(p["cpf"]):
            por_cpf[so_digitos(p["cpf"])].append(p)
    for p in produtores:
        plano.nomes[p["id"]] = p["nome_produtor"]

    produtor_da_planilha = {}  # id da planilha -> id do banco
    vistos = {}
    for r in registros:
        if r["id"]:
            vistos.setdefault(r["id"], r)
    for id_planilha, r in vistos.items():
        p = por_id.get(id_planilha)
        cpf_planilha = so_digitos(r["cpf"])
        if p is None and cpf_planilha and len(por_cpf.get(cpf_planilha, [])) == 1:
            p = por_cpf[cpf_planilha][0]
            plano.avisos.append(
                f"Produtor {id_planilha} ({r['nome'] or '?'}) nao existe com esse ID; "
                f"usei o CPF e encontrei o ID {p['id']}."
            )
        if p is None:
            plano.bloqueios.append(
                f"Produtor {id_planilha} ({r['nome'] or '?'}, linha {r['linha']}) nao existe no sistema."
            )
            continue
        cpf_banco = so_digitos(p["cpf"])
        if cpf_planilha and cpf_banco and cpf_planilha != cpf_banco:
            plano.bloqueios.append(
                f"Produtor {id_planilha} ({r['nome'] or '?'}): o CPF da planilha nao confere com o do sistema "
                f"({p['nome_produtor']}). Confira se o ID esta certo."
            )
            continue
        if r["nome"] and sem_acento_maiusculo(r["nome"]) != sem_acento_maiusculo(p["nome_produtor"]):
            plano.avisos.append(
                f"Produtor {id_planilha}: nome na planilha ('{r['nome']}') difere do sistema ('{p['nome_produtor']}')."
            )
        produtor_da_planilha[id_planilha] = p["id"]

    # 3) Animais do banco ------------------------------------------------
    animais = conn.execute(
        "SELECT id, brinco_faec, brinco_fazenda, produtor_id, status, foto_1, foto_2, foto_3 FROM animais ORDER BY id"
    ).fetchall()
    por_brinco_banco = defaultdict(list)
    for a in animais:
        por_brinco_banco[normalizar_brinco(a["brinco_faec"])].append(a)
    oficiais = _pares_oficiais()

    # 4) Movimentos --------------------------------------------------------
    destino_final = {a["id"]: a["produtor_id"] for a in animais}  # para a conferencia do limite
    animais_da_planilha = set()   # ids de animais do banco que a planilha cobre
    for (brinco, fz), usos in sorted(por_animal.items()):
        ids = {i for i, _, _ in usos}
        if len(ids) != 1 and not (tem_fazenda and fz):
            continue  # ja bloqueado acima
        id_planilha = sorted(ids)[0]
        novo_id = produtor_da_planilha.get(id_planilha)
        if novo_id is None:
            continue  # produtor com problema, ja bloqueado
        fz_original = usos[0][2]

        candidatos = por_brinco_banco.get(brinco, [])
        if tem_fazenda and fz:
            todos = candidatos
            candidatos = [c for c in candidatos if chave_fazenda(c["brinco_fazenda"]) == fz]
            if todos and not candidatos:
                atrelados = ", ".join(sorted({f"'{c['brinco_fazenda'] or 'sem brinco da fazenda'}'" for c in todos}))
                plano.bloqueios.append(
                    f"O brinco FAEC {brinco} esta atrelado, no sistema, ao brinco da fazenda {atrelados}, "
                    f"mas a planilha diz '{fz_original}'. Confira qual esta certo (o animal e a foto vao junto com o par)."
                )
                continue
        if not candidatos:
            plano.sem_cadastro.append((brinco, id_planilha))
            continue
        if len(candidatos) > 1:
            plano.bloqueios.append(
                f"O brinco {brinco} existe {len(candidatos)} vezes no banco "
                f"(ids de animal {', '.join(str(c['id']) for c in candidatos)}; fazenda: "
                f"{', '.join(repr(c['brinco_fazenda'] or '-') for c in candidatos)}); nao sei qual atualizar. "
                f"Inclua a coluna 'Brinco Fazenda' na planilha."
            )
            continue
        a = candidatos[0]
        animais_da_planilha.add(a["id"])
        if a["produtor_id"] == novo_id and a["status"] == "alocado":
            plano.iguais += 1
            continue
        plano.movimentos.append({
            "acao": "alocar" if a["produtor_id"] is None else "trocar de produtor",
            "animal_id": a["id"], "brinco": brinco, "fazenda": a["brinco_fazenda"],
            "fotos": _qtd_fotos(a),
            "atual_id": a["produtor_id"], "novo_id": novo_id, "novo_status": "alocado",
        })
        destino_final[a["id"]] = novo_id
        oficial = oficiais.get(brinco)
        if oficial and a["brinco_fazenda"] and chave_fazenda(oficial) != chave_fazenda(a["brinco_fazenda"]):
            plano.avisos.append(
                f"Brinco FAEC {brinco}: no sistema esta atrelado a '{a['brinco_fazenda']}', mas a lista oficial da "
                f"importacao inicial diz '{oficial}'. Confira se o par nao foi alterado por engano."
            )

    # 5) Alocados no banco e ausentes da planilha ---------------------------
    for a in animais:
        if a["produtor_id"] is not None and a["id"] not in animais_da_planilha:
            plano.ausentes.append(a)
            if liberar_ausentes:
                plano.movimentos.append({
                    "acao": "liberar (ausente da planilha)",
                    "animal_id": a["id"], "brinco": normalizar_brinco(a["brinco_faec"]),
                    "fazenda": a["brinco_fazenda"], "fotos": _qtd_fotos(a),
                    "atual_id": a["produtor_id"], "novo_id": None, "novo_status": "disponivel",
                })
                destino_final[a["id"]] = None

    # 6) Limite por produtor ------------------------------------------------
    contagem = Counter(pid for pid in destino_final.values() if pid is not None)
    for pid, qtd in sorted(contagem.items()):
        if qtd > LIMITE_POR_PRODUTOR:
            msg = (f"{plano.nomes.get(pid, pid)} (ID {pid}) ficaria com {qtd} animais "
                   f"(limite {LIMITE_POR_PRODUTOR}).")
            if ignorar_limite:
                plano.avisos.append(msg)
            else:
                plano.bloqueios.append(msg + " Use --ignorar-limite se for proposital.")
    return plano


# ------------------------------------------------------------------
# Gravacao protegida
# ------------------------------------------------------------------
def _impressao_das_fotos(conn):
    """Impressao digital das fotos de TODOS os animais (id + foto_1/2/3)."""
    linhas = conn.execute(
        "SELECT id, foto_1, foto_2, foto_3 FROM animais ORDER BY id"
    ).fetchall()
    h = hashlib.sha256()
    for r in linhas:
        h.update(f"{r['id']}|{r['foto_1'] or ''}|{r['foto_2'] or ''}|{r['foto_3'] or ''}\n".encode("utf-8"))
    return len(linhas), h.hexdigest()


def _salvar_backup(conn, carimbo):
    PASTA_SAIDA.mkdir(exist_ok=True)
    linhas = conn.execute("SELECT * FROM animais ORDER BY id").fetchall()
    caminho = PASTA_SAIDA / f"backup_animais_{carimbo}.csv"
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        if linhas:
            w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()))
            w.writeheader()
            for r in linhas:
                w.writerow({k: ("" if v is None else v) for k, v in r.items()})
    return caminho


def aplicar(conn, plano, carimbo):
    raw = conn._conn  # conexao do psycopg2 por tras do proxy
    backup = _salvar_backup(conn, carimbo)
    print(f"Backup da tabela de animais salvo em: {backup}")

    total_antes, fotos_antes = _impressao_das_fotos(conn)
    try:
        for m in plano.movimentos:
            cur = conn.execute(
                # UNICO comando que grava: so muda dono/status. Nunca toca nas fotos.
                "UPDATE animais SET produtor_id = ?, status = ?, atualizado_em = now() "
                "WHERE id = ? AND produtor_id IS NOT DISTINCT FROM ?",
                (m["novo_id"], m["novo_status"], m["animal_id"], m["atual_id"]),
            )
            if cur._cursor.rowcount != 1:
                raise RuntimeError(
                    f"O animal {m['brinco']} (id {m['animal_id']}) mudou no banco enquanto o script rodava. "
                    "Nada foi gravado; rode a simulacao de novo."
                )
        total_depois, fotos_depois = _impressao_das_fotos(conn)
        if total_depois != total_antes or fotos_depois != fotos_antes:
            raise RuntimeError(
                "TRAVA DE SEGURANCA: as fotos dos animais mudariam com esta operacao. "
                "Nada foi gravado."
            )
        raw.commit()
    except Exception:
        raw.rollback()
        raise
    return backup


# ------------------------------------------------------------------
# Relatorio
# ------------------------------------------------------------------
def _nome(plano, pid):
    return "(sem produtor)" if pid is None else f"{plano.nomes.get(pid, '?')} [ID {pid}]"


def imprimir_relatorio(plano, aplicado):
    print()
    print("=" * 70)
    print("RESULTADO" + (" (GRAVADO NO BANCO)" if aplicado else " - SIMULACAO (nada foi gravado)"))
    print("=" * 70)
    print(f"Animais que ja estavam certos ................ {plano.iguais}")
    print(f"Animais que MUDAM de dono .................... {len(plano.movimentos)}")
    print(f"Brincos da planilha sem cadastro no banco .... {len(plano.sem_cadastro)}")
    print(f"Vagas em aberto na planilha (sem brinco) ..... {len(plano.vagas)}")
    print(f"Alocados no banco e ausentes da planilha ..... {len(plano.ausentes)}")
    print("As fotos de todos os animais permanecem como estao.")

    if plano.movimentos:
        print("\n--- Mudancas ---")
        for m in plano.movimentos:
            print(f"  brinco {m['brinco']} (fazenda {m['fazenda'] or '-'}, {m['fotos']} foto(s) vao junto): "
                  f"{_nome(plano, m['atual_id'])}  ->  {_nome(plano, m['novo_id'])}")
    if plano.sem_cadastro:
        print("\n--- Brincos da planilha que NAO existem no banco (nada foi feito com eles) ---")
        for brinco, pid in plano.sem_cadastro:
            print(f"  brinco {brinco} (planilha: produtor {pid})")
    if plano.vagas:
        print("\n--- Vagas em aberto (produtor na planilha, mas sem brinco) ---")
        cont = Counter(r["id"] for r in plano.vagas)
        for pid, qtd in cont.items():
            nome = next((r["nome"] for r in plano.vagas if r["id"] == pid), "?")
            print(f"  produtor {pid} ({nome}): {qtd} vaga(s) sem brinco")
    if plano.ausentes:
        print("\n--- Alocados no banco, mas fora da planilha"
              + (" (SERAO LIBERADOS)" if any(m["acao"].startswith("liberar") for m in plano.movimentos) else " (nao foram tocados)") + " ---")
        for a in plano.ausentes:
            print(f"  brinco {normalizar_brinco(a['brinco_faec'])} (fazenda {a['brinco_fazenda'] or '-'}): {_nome(plano, a['produtor_id'])}")
    if plano.avisos:
        print("\n--- Avisos ---")
        for t in plano.avisos:
            print("  -", t)
    if plano.bloqueios:
        print("\n--- PROBLEMAS QUE IMPEDEM A GRAVACAO (corrija a planilha) ---")
        for t in plano.bloqueios:
            print("  X", t)


def salvar_relatorio_csv(plano, carimbo, aplicado):
    PASTA_SAIDA.mkdir(exist_ok=True)
    caminho = PASTA_SAIDA / f"{'aplicado' if aplicado else 'simulacao'}_alocacoes_{carimbo}.csv"
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["acao", "brinco_faec", "brinco_fazenda", "fotos_que_vao_junto", "animal_id", "produtor_atual", "produtor_novo"])
        for m in plano.movimentos:
            w.writerow([m["acao"], m["brinco"], m["fazenda"] or "", m["fotos"], m["animal_id"],
                        _nome(plano, m["atual_id"]), _nome(plano, m["novo_id"])])
    return caminho


# ------------------------------------------------------------------
def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="Atualiza produtor <-> brinco a partir da planilha, sem mexer nas fotos.")
    ap.add_argument("planilha", help="arquivo .xlsx (ex.: animais_alocados_fats.xlsx)")
    ap.add_argument("--aplicar", action="store_true", help="grava no banco (padrao: so simula)")
    ap.add_argument("--liberar-ausentes", action="store_true",
                    help="animais alocados no banco e ausentes da planilha voltam a 'disponivel'")
    ap.add_argument("--ignorar-limite", action="store_true", help="permite passar de 5 animais por produtor")
    args = ap.parse_args()

    caminho = Path(args.planilha)
    if not caminho.exists():
        sys.exit(f"Arquivo nao encontrado: {caminho}")

    registros, tem_fazenda = ler_planilha(caminho)
    print(f"Planilha lida: {len(registros)} linha(s) com dados.")
    if not tem_fazenda:
        print("AVISO: a planilha nao tem a coluna 'Brinco Fazenda'. Sem ela nao da para conferir se cada brinco FAEC "
              "esta atrelado ao brinco da fazenda certo. Recomendo usar o export do sistema (que traz as duas colunas).")

    conn = get_connection()
    try:
        if not args.aplicar:
            conn.execute("SET TRANSACTION READ ONLY")  # na simulacao o banco nem aceita gravar
        plano = montar_plano(conn, registros, args.liberar_ausentes, args.ignorar_limite, tem_fazenda)
        carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")

        if args.aplicar and plano.bloqueios:
            imprimir_relatorio(plano, aplicado=False)
            print("\nNADA FOI GRAVADO: corrija os problemas acima e rode de novo.")
            sys.exit(2)

        if args.aplicar and plano.movimentos:
            try:
                aplicar(conn, plano, carimbo)
            except RuntimeError as erro:
                print(f"\nERRO: {erro}")
                print("O banco foi deixado exatamente como estava.")
                sys.exit(1)
        elif args.aplicar:
            print("Nada para alterar: o banco ja esta igual a planilha.")

        imprimir_relatorio(plano, aplicado=args.aplicar and bool(plano.movimentos))
        if plano.movimentos:
            rel = salvar_relatorio_csv(plano, carimbo, args.aplicar)
            print(f"\nRelatorio salvo em: {rel}")
        if not args.aplicar:
            if plano.bloqueios:
                print("\nCorrija os problemas acima antes de usar --aplicar.")
                sys.exit(2)
            print("\nSe estiver certo, rode de novo com --aplicar para gravar.")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
