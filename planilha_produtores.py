"""
Importação de produtores a partir de uma planilha do Excel (.xlsx).

Colunas reconhecidas (pelo nome do cabeçalho, em qualquer ordem):
    Nome, CPF, RG, Propriedade, Município, Telefone

A planilha pode ter linhas de título ou linhas em branco antes do
cabeçalho (como "PROGRAMA MATRIZES DA MANHÃ"): o cabeçalho é procurado
automaticamente nas primeiras linhas.

Como ligar no sistema (app.py) — duas linhas, antes do
`if __name__ == "__main__":`:

    from planilha_produtores import registrar_importacao
    registrar_importacao(app, get_connection, _formatar_cpf, _formatar_telefone, _erro_de_conexao)
"""

import io
import re
import unicodedata
from datetime import datetime
from zoneinfo import ZoneInfo

from campos import FIELDS

# nome do campo -> variações aceitas no cabeçalho (já sem acento e minúsculas)
_CABECALHOS = {
    "nome": {"nome", "produtor", "nome do produtor", "nome produtor"},
    "cpf": {"cpf"},
    "rg": {"rg"},
    "propriedade": {"propriedade", "nome da propriedade", "fazenda"},
    "municipio": {"municipio", "cidade"},
    "telefone": {"telefone", "celular", "fone", "contato"},
}
_MAX_LINHAS_PROCURANDO_CABECALHO = 20


def _sem_acento(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in texto if not unicodedata.combining(c))


def _chave_cabecalho(valor):
    if valor is None:
        return None
    limpo = re.sub(r"\s+", " ", _sem_acento(valor).strip().lower())
    for campo, variacoes in _CABECALHOS.items():
        if limpo in variacoes:
            return campo
    return None


def _como_texto(valor, campo):
    """Converte a célula em texto. Números (RG, telefone, CPF sem formatação)
    chegam do Excel como int/float, e CPF perde os zeros à esquerda."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if isinstance(valor, int):
        texto = str(valor)
        return texto.zfill(11) if campo == "cpf" else texto
    return str(valor).strip()


def ler_planilha_produtores(conteudo_xlsx):
    """Lê o .xlsx e devolve (linhas, erro).

    `linhas` é uma lista de dicts com as chaves: linha (número na planilha),
    nome, cpf, rg, propriedade, municipio, telefone — todos em texto puro,
    sem nenhuma formatação aplicada. `erro` é uma mensagem (ou None)."""
    from openpyxl import load_workbook

    try:
        planilha = load_workbook(io.BytesIO(conteudo_xlsx), data_only=True, read_only=True).active
        todas = list(planilha.iter_rows(values_only=True))
    except Exception:
        return [], "Não foi possível ler o arquivo. Confira se é uma planilha do Excel (.xlsx)."

    mapa_colunas = None
    linha_cabecalho = None
    for i, linha in enumerate(todas[:_MAX_LINHAS_PROCURANDO_CABECALHO]):
        encontrados = {}
        for j, celula in enumerate(linha):
            campo = _chave_cabecalho(celula)
            if campo and campo not in encontrados:
                encontrados[campo] = j
        if "nome" in encontrados and "cpf" in encontrados:
            mapa_colunas, linha_cabecalho = encontrados, i
            break

    if mapa_colunas is None:
        return [], (
            "Não encontrei o cabeçalho da planilha. Ela precisa ter, no mínimo, "
            "as colunas \"Nome\" e \"CPF\" (as outras são: RG, Propriedade, Município e Telefone)."
        )

    linhas = []
    for i, linha in enumerate(todas[linha_cabecalho + 1:], start=linha_cabecalho + 2):
        registro = {"linha": i}
        for campo in _CABECALHOS:
            j = mapa_colunas.get(campo)
            registro[campo] = _como_texto(linha[j], campo) if j is not None and j < len(linha) else ""
        if not any(registro[c] for c in _CABECALHOS):
            continue  # linha totalmente em branco
        linhas.append(registro)
    return linhas, None


def _resumo_nomes(nomes, limite=8):
    if len(nomes) <= limite:
        return ", ".join(nomes)
    return ", ".join(nomes[:limite]) + f" e mais {len(nomes) - limite}"


def registrar_importacao(app, get_connection, formatar_cpf, formatar_telefone, erro_de_conexao):
    """Cria a rota POST /produtores/importar (endpoint `importar_produtores`)."""
    from flask import flash, redirect, request, url_for

    def importar_produtores():
        arquivo = request.files.get("planilha")
        if not arquivo or not arquivo.filename:
            flash("Escolha a planilha (.xlsx) que você quer importar.", "error")
            return redirect(url_for("index"))
        if not arquivo.filename.lower().endswith(".xlsx"):
            flash("O arquivo precisa ser uma planilha do Excel (.xlsx).", "error")
            return redirect(url_for("index"))

        linhas, erro = ler_planilha_produtores(arquivo.read())
        if erro:
            flash(erro, "error")
            return redirect(url_for("index"))
        if not linhas:
            flash("A planilha não tem nenhum produtor para importar.", "error")
            return redirect(url_for("index"))

        importados, atualizados, repetidos, invalidos = [], [], [], []
        conn = None
        try:
            conn = get_connection()
            # CPF (só dígitos) -> id e RG atual de quem já está cadastrado
            existentes = {}
            for r in conn.execute("SELECT id, cpf, rg FROM produtores ORDER BY id").fetchall():
                chave = re.sub(r"\D", "", r["cpf"] or "")
                if chave and chave not in existentes:
                    existentes[chave] = {"id": r["id"], "rg": (r["rg"] or "").strip()}
            hoje = datetime.now(ZoneInfo("America/Fortaleza")).strftime("%d/%m/%Y")

            for l in linhas:
                nome = l["nome"].strip().upper()
                cpf_digitos = re.sub(r"\D", "", l["cpf"])
                rotulo = nome or f"linha {l['linha']}"

                if not nome:
                    invalidos.append(f"linha {l['linha']} (sem nome)")
                    continue
                if len(cpf_digitos) != 11:
                    invalidos.append(f"{nome} (CPF ausente ou inválido)")
                    continue
                if cpf_digitos in existentes:
                    # Já cadastrado: não duplica. Só preenche o RG se estiver
                    # vazio no sistema e vier preenchido na planilha (nunca
                    # sobrescreve um RG que já existe).
                    atual = existentes[cpf_digitos]
                    rg_planilha = l["rg"].strip()
                    if rg_planilha and not atual["rg"]:
                        conn.execute(
                            "UPDATE produtores SET rg = ?, atualizado_em = now() WHERE id = ?",
                            (rg_planilha, atual["id"]),
                        )
                        atual["rg"] = rg_planilha
                        atualizados.append(nome)
                    else:
                        repetidos.append(nome)
                    continue

                dados = {campo: "" for campo in FIELDS}
                dados["nome_produtor"] = nome
                dados["cpf"] = formatar_cpf(cpf_digitos)
                dados["rg"] = l["rg"].strip()
                dados["telefone"] = formatar_telefone(l["telefone"])
                dados["nome_propriedade"] = l["propriedade"].strip().upper()
                dados["municipio"] = l["municipio"].strip().upper()

                novo = conn.execute(
                    f"INSERT INTO produtores ({', '.join(FIELDS)}) "
                    f"VALUES ({', '.join(['?'] * len(FIELDS))}) RETURNING id",
                    [dados[campo] for campo in FIELDS],
                ).fetchone()
                # Mesma "visita inicial" que o cadastro manual (/novo) cria.
                conn.execute(
                    "INSERT INTO visitas "
                    "(produtor_id, data_visita, tipo_visita, observacoes, etapa, ativa) "
                    "VALUES (?, ?, ?, ?, 1, TRUE)",
                    (novo["id"], hoje, "Cadastro / Ficha inicial",
                     "Realizar cadastro inicial dos produtores"),
                )
                existentes[cpf_digitos] = {"id": novo["id"], "rg": dados["rg"]}
                importados.append(nome)

            conn.commit()  # tudo ou nada: só grava se chegou até aqui sem erro
        except Exception as e:
            if erro_de_conexao(e):
                flash(
                    "Sem conexão com a internet no momento — a importação precisa "
                    "de internet. Nada foi importado; tente de novo quando voltar.",
                    "error",
                )
            else:
                app.logger.exception("Falha ao importar planilha de produtores")
                flash(f"Não foi possível importar a planilha (nada foi gravado): {e}", "error")
            return redirect(url_for("index"))
        finally:
            if conn is not None:
                conn.close()

        if importados:
            flash(f"{len(importados)} produtor(es) importado(s) com sucesso.", "success")
        elif not atualizados:
            flash("Nenhum produtor novo foi importado.", "aviso")
        if atualizados:
            flash(
                f"{len(atualizados)} produtor(es) já cadastrado(s) tiveram o RG preenchido: "
                f"{_resumo_nomes(atualizados)}.",
                "success",
            )
        if repetidos:
            flash(
                f"{len(repetidos)} já estavam cadastrados (mesmo CPF), sem nada novo para preencher: "
                f"{_resumo_nomes(repetidos)}.",
                "aviso",
            )
        if invalidos:
            flash(
                f"{len(invalidos)} linha(s) não foram importadas: {_resumo_nomes(invalidos)}.",
                "error",
            )
        return redirect(url_for("index"))

    app.add_url_rule(
        "/produtores/importar", "importar_produtores", importar_produtores, methods=["POST"]
    )
