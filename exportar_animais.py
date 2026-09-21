# -*- coding: utf-8 -*-
"""
Exportação Excel dos animais alocados a cada produtor.

Módulo independente: não mexe em nenhuma rota que já existe no app.py.
Para ligar no sistema, basta 2 linhas no app.py (logo antes do
`if __name__ == "__main__":`):

    from exportar_animais import registrar_exportacao_animais
    registrar_exportacao_animais(app)

Isso cria a rota /exportar-animais-alocados (endpoint
'exportar_animais_alocados'), usada pelo botão da tela inicial.
"""

import io
import re

from flask import request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import get_connection

LIMITE_PADRAO = 5  # limite de animais por produtor no sistema


def _formatar_cpf(valor):
    digitos = re.sub(r"\D", "", valor or "")
    if len(digitos) == 11:
        return f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}"
    return valor or ""


def _filtros_produtores(busca, municipio, ateg):
    """Mesmos filtros da listagem inicial (busca, município e ATEG).
    Devolve (where_sql, params)."""
    condicoes, params = [], []
    if busca:
        like = f"%{busca}%"
        condicoes.append(
            "(p.nome_produtor ILIKE ? OR p.cpf ILIKE ? OR p.municipio ILIKE ? "
            "OR p.nome_propriedade ILIKE ?)"
        )
        params += [like, like, like, like]
    if municipio:
        condicoes.append("p.municipio = ?")
        params.append(municipio)
    if ateg in ("Sim", "Não"):
        condicoes.append("p.assistido_ateg = ?")
        params.append(ateg)
    return ("WHERE " + " AND ".join(condicoes)) if condicoes else "", params


def _peso_inteiro(valor):
    """O peso é guardado como texto (ex: '320,5'). Vira número INTEIRO
    (arredondado, 320,5 -> 321) quando der."""
    if valor is None or str(valor).strip() == "":
        return None
    try:
        return int(float(str(valor).strip().replace(",", ".")) + 0.5)
    except ValueError:
        return str(valor).strip()


def montar_excel(linhas, meta=LIMITE_PADRAO):
    """`linhas` vem de um LEFT JOIN produtores x animais (produtor sem
    animal vem com brinco_faec vazio). Devolve um BytesIO com o .xlsx."""
    wb = Workbook()
    fill = PatternFill(start_color="1F4E2C", end_color="1F4E2C", fill_type="solid")
    fonte = Font(color="FFFFFF", bold=True)

    def cabecalho(ws, titulos):
        ws.append(titulos)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = fonte
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

    # ---- Aba 1: um animal por linha, com o produtor ao qual foi alocado ----
    ws = wb.active
    ws.title = "Animais alocados"
    cabecalho(ws, ["ID Produtor", "Produtor", "CPF", "Município", "Propriedade",
                   "Técnico", "ATEG", "Brinco FAEC", "Brinco Fazenda", "Peso (kg)"])
    n_animais = 0
    for r in linhas:
        if not r["brinco_faec"]:
            continue  # produtor sem animal: só aparece no Resumo
        ws.append([
            r["produtor_id"], (r["nome_produtor"] or "").upper(),
            _formatar_cpf(r["cpf"]), r["municipio"], r["nome_propriedade"],
            r["tecnico_responsavel"], r["assistido_ateg"],
            r["brinco_faec"], r["brinco_fazenda"], _peso_inteiro(r["peso"]),
        ])
        n_animais += 1
        linha = ws.max_row
        for col in (3, 8, 9):  # CPF e brincos como texto (mantém zeros à esquerda)
            ws.cell(row=linha, column=col).number_format = "@"
        ws.cell(row=linha, column=10).number_format = "0"
    ultima = max(n_animais + 1, 2)
    for i, larg in enumerate([11, 38, 16, 20, 28, 18, 8, 13, 15, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    if n_animais:
        ws.auto_filter.ref = f"A1:J{ultima}"

    # ---- Aba 2: resumo por produtor (todos os produtores do filtro) ----
    rs = wb.create_sheet("Resumo por produtor")
    cabecalho(rs, ["ID", "Produtor", "Município", "Propriedade",
                   "Animais alocados", "Faltam p/ meta"])
    rs["H1"] = "Meta por produtor"
    rs["H1"].font = Font(bold=True)
    rs["I1"] = meta
    rs["I1"].font = Font(color="0000FF", bold=True)  # valor de entrada (pode editar)
    rs["J1"] = "Limite de animais por produtor no sistema (editável)."
    rs["J1"].font = Font(italic=True, size=9, color="7F7F7F")

    vistos = set()
    for r in linhas:
        if r["produtor_id"] in vistos:
            continue
        vistos.add(r["produtor_id"])
        rs.append([r["produtor_id"], (r["nome_produtor"] or "").upper(),
                   r["municipio"], r["nome_propriedade"]])
        i = rs.max_row
        rs.cell(row=i, column=5,
                value=f"=COUNTIF('Animais alocados'!$A$2:$A${ultima},A{i})")
        rs.cell(row=i, column=6, value=f"=MAX(0,$I$1-E{i})")
    fim = rs.max_row
    if vistos:
        rs.cell(row=fim + 1, column=4, value="Total").font = Font(bold=True)
        c = rs.cell(row=fim + 1, column=5, value=f"=SUM(E2:E{fim})")
        c.font = Font(bold=True)
        rs.auto_filter.ref = f"A1:F{fim}"
    for i, larg in enumerate([8, 38, 20, 28, 17, 15, 3, 18, 8, 46], start=1):
        rs.column_dimensions[get_column_letter(i)].width = larg

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def exportar_animais_alocados():
    """Excel com os animais alocados a cada produtor. Respeita os mesmos
    filtros da tela (busca, município e ATEG)."""
    where_sql, params = _filtros_produtores(
        request.args.get("q", "").strip(),
        request.args.get("municipio", "").strip(),
        request.args.get("ateg", "").strip(),
    )
    conn = get_connection()
    try:
        rows = conn.execute(
            f"""SELECT p.id AS produtor_id, p.nome_produtor, p.cpf, p.municipio,
                       p.nome_propriedade, p.tecnico_responsavel, p.assistido_ateg,
                       a.brinco_faec, a.brinco_fazenda, a.peso
                FROM produtores p
                LEFT JOIN animais a ON a.produtor_id = p.id
                {where_sql}
                ORDER BY p.nome_produtor, a.brinco_faec""",
            params,
        ).fetchall()
    finally:
        conn.close()

    return send_file(
        montar_excel([dict(r) for r in rows]),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="animais_alocados_fats.xlsx",
    )


def registrar_exportacao_animais(app):
    app.add_url_rule(
        "/exportar-animais-alocados",
        endpoint="exportar_animais_alocados",
        view_func=exportar_animais_alocados,
    )
