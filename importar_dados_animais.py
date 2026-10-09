"""
Importação dos DADOS INDIVIDUAIS dos animais a partir de uma planilha do Excel (.xlsx).

A planilha atualiza animais que JÁ existem no sistema (o animal é achado pelo
Brinco FAEC). Colunas reconhecidas (pelo nome do cabeçalho, em qualquer ordem):
    Brinco FAEC       (obrigatória)
    Brinco Fazenda    (opcional)  -> só CONFERE com o animal; se for diferente, a linha não é importada
    Peso Atual / Peso (opcional)  -> preenche o peso SÓ de quem está sem peso (não troca peso existente)
    Grau de Sangue    (opcional)
    Pai               (opcional)
    Brinco da Mãe     (opcional)
    Data de Nascimento(opcional)

Regras:
  - célula vazia na planilha NUNCA apaga o que já está no sistema;
  - nada é gravado na primeira etapa (prévia); a confirmação revalida tudo;
  - tudo ou nada: se der erro na gravação, nada é alterado.

Como ligar no sistema (app.py), antes do `if __name__ == "__main__":`:

    from importar_dados_animais import registrar_importacao_dados_animais
    registrar_importacao_dados_animais(app, get_connection, _erro_de_conexao)
"""

import io
import json
import re
import unicodedata
from datetime import date, datetime

_MAX_LINHAS_PROCURANDO_CABECALHO = 20

# campo -> variações aceitas no cabeçalho (sem acento, minúsculas)
_CABECALHOS = {
    "brinco_faec": {"brinco faec", "brinco_faec", "brinco"},
    "brinco_fazenda": {"brinco fazenda", "brinco da fazenda", "brinco_fazenda"},
    "peso": {"peso atual", "peso", "peso (kg)", "peso kg"},
    "grau_sangue": {"grau de sangue", "grau sangue", "grau_sangue"},
    "pai": {"pai", "touro pai"},
    "brinco_mae": {"brinco da mae", "brinco mae", "brinco_mae", "mae"},
    "data_nascimento": {"data de nascimento", "data nascimento", "nascimento", "data_nascimento"},
}
ROTULOS = {
    "peso": "Peso (kg)", "grau_sangue": "Grau de sangue", "pai": "Pai",
    "brinco_mae": "Brinco da mãe", "data_nascimento": "Nascimento",
}
# Brincos NÃO são alterados (servem só para achar e conferir o animal).
# O peso só é preenchido quando o animal ainda está sem peso; nunca troca um peso existente.
CAMPOS_ATUALIZAVEIS = ["peso", "grau_sangue", "pai", "brinco_mae", "data_nascimento"]
SO_SE_VAZIO = {"peso"}


def _sem_acento(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in texto if not unicodedata.combining(c))


def _chave_cabecalho(valor):
    if valor is None:
        return None
    limpo = re.sub(r"\s+", " ", _sem_acento(valor).strip().lower())
    for campo, variacoes in _CABECALHOS.items():
        if limpo in variacoes:
            return campo
    return None


def chave_brinco_faec(valor):
    """'0052', 52, 52.0 e ' 52 ' viram todos '52'; 'V1694' continua 'V1694'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = re.sub(r"[\s\-./]", "", str(valor)).upper()
    return str(int(texto)) if texto.isdigit() else texto


def chave_fazenda(valor):
    return re.sub(r"[\s\-./]", "", str(valor or "")).upper()


def _texto_maiusculo(valor):
    return " ".join(str(valor or "").split()).upper()


def _peso_texto(valor):
    """255 / 255.0 / '255,5' -> '255' / '255' / '255,5'. Devolve '' se não for número."""
    if valor is None or str(valor).strip() == "":
        return ""
    if isinstance(valor, (int, float)):
        numero = float(valor)
    else:
        try:
            numero = float(str(valor).strip().replace(".", "").replace(",", ".") if "," in str(valor) else str(valor).strip())
        except ValueError:
            return ""
    if numero <= 0 or numero > 3000:
        return ""
    return str(int(numero)) if numero.is_integer() else f"{numero:.1f}".replace(".", ",")


def _peso_igual(a, b):
    def num(x):
        try:
            return float(str(x).replace(",", "."))
        except ValueError:
            return None
    return num(a) is not None and num(a) == num(b)


def _data_texto(valor):
    """datetime/date/'dd/mm/aaaa'/'aaaa-mm-dd' -> 'dd/mm/aaaa'. '' se inválida."""
    if valor is None or str(valor).strip() == "":
        return ""
    if isinstance(valor, (datetime, date)):
        d = valor
    else:
        texto = str(valor).strip()
        d = None
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
            try:
                d = datetime.strptime(texto[:10] if formato == "%Y-%m-%d" else texto, formato)
                break
            except ValueError:
                continue
        if d is None:
            return ""
    if d.year < 1990 or d.year > date.today().year:
        return ""
    return f"{d.day:02d}/{d.month:02d}/{d.year}"


def ler_planilha_dados(conteudo_xlsx):
    """Lê o .xlsx e devolve (linhas, erro). Cada linha é um dict já normalizado:
    {linha, brinco_faec (como veio), brinco_fazenda, peso, grau_sangue, pai,
     brinco_mae, data_nascimento, invalidos: [campos com valor inválido]}."""
    from openpyxl import load_workbook

    try:
        planilha = load_workbook(io.BytesIO(conteudo_xlsx), data_only=True, read_only=True).worksheets[0]
        todas = list(planilha.iter_rows(values_only=True))
    except Exception:
        return [], "Não foi possível ler o arquivo. Confira se é uma planilha do Excel (.xlsx)."

    mapa, linha_cab = None, None
    for i, linha in enumerate(todas[:_MAX_LINHAS_PROCURANDO_CABECALHO]):
        achados = {}
        for j, celula in enumerate(linha):
            campo = _chave_cabecalho(celula)
            if campo and campo not in achados:
                achados[campo] = j
        if "brinco_faec" in achados and any(c in achados for c in CAMPOS_ATUALIZAVEIS):
            mapa, linha_cab = achados, i
            break
    if mapa is None:
        return [], (
            "Não encontrei o cabeçalho da planilha. Ela precisa ter a coluna \"Brinco FAEC\" e ao menos uma de: "
            "Peso Atual, Grau de Sangue, Pai, Brinco da Mãe, Data de Nascimento."
        )

    def celula(linha, campo):
        j = mapa.get(campo)
        return linha[j] if j is not None and j < len(linha) else None

    linhas = []
    for i, linha in enumerate(todas[linha_cab + 1:], start=linha_cab + 2):
        bruto = celula(linha, "brinco_faec")
        if bruto is None or str(bruto).strip() == "":
            continue
        reg = {"linha": i, "brinco_faec": str(bruto if not isinstance(bruto, float) else int(bruto)).strip(),
               "brinco_fazenda": str(celula(linha, "brinco_fazenda") or "").strip(), "invalidos": []}
        for campo, converter in (("peso", _peso_texto), ("data_nascimento", _data_texto)):
            bruto_campo = celula(linha, campo)
            valor = converter(bruto_campo)
            if bruto_campo not in (None, "") and not valor:
                reg["invalidos"].append(campo)
            reg[campo] = valor
        for campo in ("grau_sangue", "pai", "brinco_mae"):
            reg[campo] = _texto_maiusculo(celula(linha, campo))
        linhas.append(reg)
    return linhas, None


def validar_dados(linhas, animais):
    """Lógica pura (sem banco). `animais` = lista de dicts com id, brinco_faec,
    brinco_fazenda, peso, grau_sangue, pai, brinco_mae, data_nascimento.
    Devolve {"atualizar": [...], "sem_mudanca": [...], "problemas": [...]}."""
    por_brinco = {}
    for a in animais:
        por_brinco.setdefault(chave_brinco_faec(a["brinco_faec"]), []).append(a)

    problemas, atualizar, sem_mudanca, vistos = [], [], [], {}

    def erro(l, motivo):
        problemas.append({"linha": l["linha"], "brinco": l["brinco_faec"], "motivo": motivo})

    for l in linhas:
        chave = chave_brinco_faec(l["brinco_faec"])
        if chave in vistos:
            erro(l, f"Brinco FAEC repetido na planilha (já na linha {vistos[chave]})")
            continue
        vistos[chave] = l["linha"]
        if l.get("invalidos"):
            nomes = ", ".join(ROTULOS[c] for c in l["invalidos"])
            erro(l, f"valor inválido em: {nomes}")
            continue
        achados = por_brinco.get(chave, [])
        if not achados:
            erro(l, "brinco não cadastrado em Animais")
            continue
        if len(achados) > 1:
            erro(l, f"{len(achados)} animais com esse Brinco FAEC (ambíguo)")
            continue
        a = achados[0]
        if l["brinco_fazenda"] and a["brinco_fazenda"] and chave_fazenda(l["brinco_fazenda"]) != chave_fazenda(a["brinco_fazenda"]):
            erro(l, f"o Brinco Fazenda da planilha ({l['brinco_fazenda']}) é diferente do cadastrado ({a['brinco_fazenda']}) — confira se é o mesmo animal")
            continue
        mudancas = []
        for campo in CAMPOS_ATUALIZAVEIS:
            novo = l.get(campo) or ""
            if not novo:
                continue
            antes = a.get(campo) or ""
            if campo in SO_SE_VAZIO and antes.strip():
                continue          # já tem valor: não troca
            igual = antes.strip() == novo if campo == "data_nascimento" else _texto_maiusculo(antes) == novo
            if not igual:
                mudancas.append({"campo": campo, "rotulo": ROTULOS[campo], "antes": antes, "depois": novo})
        item = {"linha": l["linha"], "brinco": a["brinco_faec"], "fazenda": a["brinco_fazenda"] or "",
                "animal_id": a["id"], "mudancas": mudancas}
        (atualizar if mudancas else sem_mudanca).append(item)
    problemas.sort(key=lambda p: p["linha"])
    return {"atualizar": atualizar, "sem_mudanca": sem_mudanca, "problemas": problemas}


def _carregar_animais(conn):
    return [dict(r) for r in conn.execute(
        "SELECT id, brinco_faec, brinco_fazenda, peso, grau_sangue, pai, brinco_mae, data_nascimento FROM animais"
    ).fetchall()]


def registrar_importacao_dados_animais(app, get_connection, erro_de_conexao):
    """Cria as rotas /animais/importar-dados (+ /confirmar)."""
    from flask import flash, redirect, render_template, request, url_for

    def _sem_conexao():
        flash("Sem conexão com a internet no momento — a importação precisa de internet. "
              "Nada foi alterado; tente de novo quando voltar.", "error")
        return redirect(url_for("importar_dados_animais"))

    def importar_dados_animais():
        if request.method == "GET":
            return render_template("importar_dados_animais.html", previa=None)
        arquivo = request.files.get("planilha")
        if not arquivo or not arquivo.filename:
            flash("Escolha a planilha (.xlsx) com os dados dos animais.", "error")
            return redirect(url_for("importar_dados_animais"))
        if not arquivo.filename.lower().endswith(".xlsx"):
            flash("O arquivo precisa ser uma planilha do Excel (.xlsx).", "error")
            return redirect(url_for("importar_dados_animais"))
        linhas, erro_leitura = ler_planilha_dados(arquivo.read())
        if erro_leitura:
            flash(erro_leitura, "error")
            return redirect(url_for("importar_dados_animais"))
        if not linhas:
            flash("A planilha não tem nenhuma linha com Brinco FAEC.", "error")
            return redirect(url_for("importar_dados_animais"))
        conn = None
        try:
            conn = get_connection()
            animais = _carregar_animais(conn)
        except Exception as e:
            if erro_de_conexao(e):
                return _sem_conexao()
            app.logger.exception("Falha ao validar planilha de dados dos animais")
            flash(f"Não foi possível validar a planilha: {e}", "error")
            return redirect(url_for("importar_dados_animais"))
        finally:
            if conn is not None:
                conn.close()
        resultado = validar_dados(linhas, animais)
        return render_template("importar_dados_animais.html", previa=resultado, linhas_json=json.dumps(linhas))

    def confirmar_importar_dados_animais():
        try:
            linhas = json.loads(request.form.get("linhas") or "[]")
        except ValueError:
            linhas = []
        if not linhas:
            flash("Nada para importar. Envie a planilha de novo.", "error")
            return redirect(url_for("importar_dados_animais"))
        conn = None
        try:
            conn = get_connection()
            resultado = validar_dados(linhas, _carregar_animais(conn))  # estado atual, não o da prévia
            for item in resultado["atualizar"]:
                sets = ", ".join(f"{m['campo']} = ?" for m in item["mudancas"])
                conn.execute(
                    f"UPDATE animais SET {sets}, atualizado_em = now() WHERE id = ?",
                    [m["depois"] for m in item["mudancas"]] + [item["animal_id"]],
                )
            conn.commit()  # tudo ou nada
        except Exception as e:
            if conn is not None:
                try:
                    conn.rollback()
                except Exception:
                    pass
            if erro_de_conexao(e):
                return _sem_conexao()
            app.logger.exception("Falha ao importar dados dos animais")
            flash("Não foi possível gravar. Nada foi alterado.", "error")
            return redirect(url_for("importar_dados_animais"))
        finally:
            if conn is not None:
                conn.close()
        flash(f"{len(resultado['atualizar'])} animal(is) atualizado(s). "
              f"{len(resultado['sem_mudanca'])} já estavam iguais.", "success")
        return redirect(url_for("animais"))

    app.add_url_rule("/animais/importar-dados", "importar_dados_animais", importar_dados_animais, methods=["GET", "POST"])
    app.add_url_rule("/animais/importar-dados/confirmar", "confirmar_importar_dados_animais",
                     confirmar_importar_dados_animais, methods=["POST"])
